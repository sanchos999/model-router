#!/usr/bin/env python3
"""R13: DOCX + PDF generator for the final Model Router manual.

Reads docs/RUKOVODSTVO-R13.md, embeds fresh screenshots from
state/e2e-shots (r12b-*/r13-*), writes:
  docs/Model_Router_v1_0_Rukovodstvo_FINAL.docx
  docs/Model_Router_v1_0_Rukovodstvo_FINAL.pdf
"""
import os
import re
import sys

from docx import Document
from docx.shared import Cm, Pt

ROOT = "/home/sanchos/model-router"
MD = os.path.join(ROOT, "docs/RUKOVODSTVO-R13.md")
SHOTS = os.path.join(ROOT, "state/e2e-shots")
OUT_DOCX = os.path.join(ROOT, "docs/Model_Router_v1_0_Rukovodstvo_FINAL.docx")
OUT_PDF = os.path.join(ROOT, "docs/Model_Router_v1_0_Rukovodstvo_FINAL.pdf")

# section -> screenshot mapping (fresh R12-B/R13 shots only)
SHOT_MAP = {
    "4. Обзор (Dashboard)": "r12b-dashboard.png",
    "5. Провайдеры": "r13-providers.png",
    "7. Каталог моделей": "r13-models.png",
    "9. Несопоставленные модели": "r13-unmatched.png",
    "10. Добавление модели": "r13-wizard-step1.png",
    "16. Price history": "r12b-tab-history.png",
    "18. Availability checks": "r12b-tab-providers.png",
    "20. Per-model minimum discount": "r12b-tab-policy.png",
    "23. Cost calculator": "r13-cost-calc.png",
    "28. Audit / Revisions": "r13-audit.png",
    "29. Known-good baseline": "r13-settings-baselines.png",
    "14. Official / market / actual pricing": "r12b-tab-prices.png",
    "12. Canonical mappings": "r13-model-detail.png",
}


def add_shots(doc, section_title: str) -> bool:
    fn = SHOT_MAP.get(section_title)
    if not fn:
        return False
    path = os.path.join(SHOTS, fn)
    if not os.path.exists(path):
        return False
    doc.add_picture(path, width=Cm(16))
    doc.paragraphs[-1].alignment = 1
    cap = doc.add_paragraph(f"Рисунок: {fn}")
    cap.alignment = 1
    for r in cap.runs:
        r.font.size = Pt(8)
        r.font.italic = True
    return True


def build_docx() -> int:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "DejaVu Sans"
    st.font.size = Pt(10)

    title = doc.add_heading("Model Router v1.0 — Руководство администратора", 0)
    title.alignment = 1
    sub = doc.add_paragraph("Финальная редакция R12-B / R13 · 12.09.2026 · коммит ff75cef")
    sub.alignment = 1

    n_shots = 0
    cur_section = ""
    for line in open(MD, encoding="utf-8").read().splitlines():
        if not line.strip():
            continue
        if line.startswith("## "):
            cur_section = line[3:].strip()
            doc.add_heading(cur_section, level=1)
            if add_shots(doc, cur_section):
                n_shots += 1
        elif line.startswith("    "):
            p = doc.add_paragraph(line.strip())
            for r in p.runs:
                r.font.name = "DejaVu Sans Mono"
                r.font.size = Pt(9)
        elif re.match(r"^\d+\.\s", line.strip()) or line.startswith("- "):
            doc.add_paragraph(line.strip(), style="List Bullet"
                              if line.startswith("- ") else "List Number")
        else:
            doc.add_paragraph(line.strip())
    doc.save(OUT_DOCX)
    return n_shots


def build_pdf() -> int:
    """PDF via PyMuPDF Story: text from the MD + embedded screenshots.
    Images referenced RELATIVE to the pymupdf.Archive dir (skill rule)."""
    import pymupdf

    html = ['<h1 style="text-align:center">Model Router v1.0 — '
            'Руководство администратора</h1>'
            '<p style="text-align:center">Финальная редакция R12-B / R13 · '
            '12.09.2026 · коммит ff75cef</p>']
    n_shots = 0
    for line in open(MD, encoding="utf-8").read().splitlines():
        if not line.strip():
            continue
        esc = (line.strip().replace("&", "&amp;").replace("<", "&lt;")
               .replace(">", "&gt;"))
        if line.startswith("## "):
            sec = line[3:].strip()
            html.append(f"<h2>{esc[3:]}</h2>")
            fn = SHOT_MAP.get(sec)
            if fn and os.path.exists(os.path.join(SHOTS, fn)):
                html.append(f'<p><img src="{fn}" style="width:400px"/></p>')
                n_shots += 1
        elif line.startswith("    "):
            html.append(f"<pre>{esc.strip()}</pre>")
        else:
            html.append(f"<p>{esc}</p>")
    story = pymupdf.Story("".join(html), archive=pymupdf.Archive(SHOTS))
    writer = pymupdf.DocumentWriter(OUT_PDF)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)
    more = 1
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    d = pymupdf.open(OUT_PDF)
    return d.page_count


if __name__ == "__main__":
    n = build_docx()
    print(f"DOCX: {OUT_DOCX} ({os.path.getsize(OUT_DOCX)} bytes, {n} screenshots)")
    pages = build_pdf()
    print(f"PDF: {OUT_PDF} pages={pages} size={os.path.getsize(OUT_PDF)}")
    sys.exit(0 if pages and pages > 0 and n >= 10 else 1)
